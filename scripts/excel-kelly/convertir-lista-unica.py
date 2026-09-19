# Convierte los Excel de Kelly a la lista única de categorías (19-sep-2026).
# Uso: exportar las reglas con scripts/excel-kelly/exportar-reglas.ts y ajustar
# las rutas S, D, OUT y el diccionario SEDES (archivo, hoja del mes, hoja de
# ventas, título, correcciones a mano por fila). Después abrir y guardar cada
# archivo en Microsoft Excel para que calcule las fórmulas.
import json, copy, sys, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter

S = "/private/tmp/claude-501/-Users-jahnnkarlo-Proyectos-Claude/9f0f60a3-7de5-4b34-a1ef-0245f228f274/scratchpad/"
D = "/Users/jahnnkarlo/Library/CloudStorage/Dropbox/1. Yayi's/01. Grupo Yayi's/02. Finanzas/Reportes Kelly/Semana 38/"
OUT = S + "excel-nuevo/"
R = json.load(open(S + "reglas.json"))
REGLAS, CATS = R["reglas"], R["categorias"]

SEDES = {
  "ATELIER": ("INGRESOS_GASTOS_SUMMARY_ATELIER_15_09_2026_PE_AUTOMATICO.xlsx", "Ing&Gtos SEP26", "Control de VTAS-SEP", "YAYI'S ATELIER", {123: "AHORRO"}),
  "FONAVI": ("INGRESOS_GASTOS_SUMMARY_FONAVI_15_09_2026_PE_AUTOMATICO.xlsx", "Ing&Gtos SEP26", "Control de VTAS-SEP", "YAYI'S FONAVI", {29: "MENAJE Y UTENSILIOS"}),
  "CENTRO": ("INGRESOS_GASTOS_SUMMARY_CENTRO_15_09_2026_PE_AUTOMATICO.xlsx", "Ing&Gtos-SEP26", "Control de VTAS-SEP", "YAYI'S CENTRO", {}),
}

VERDE, OSCURO, CREMA = "098B5F", "004C40", "F9F6EB"
H = Font(name="Calibri", bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor=OSCURO)
TIPO_FILL = {"Fijo": "E3F1EA", "Variable": "FFF4D6", "Financiamiento": "E8E4F4", "Inversión": "E1ECF7", "No es gasto": "EFEFEF"}
thin = Side(style="thin", color="D9D9D9")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

def acentos(expr):
    for a, b in [("Á","A"),("É","E"),("Í","I"),("Ó","O"),("Ú","U"),("À","A"),("È","E"),("Ì","I"),("Ò","O"),("Ù","U"),("Ñ","N"),("Ü","U")]:
        expr = f'SUBSTITUTE({expr},"{a}","{b}")'
    return expr

def hoja_catalogo(wb):
    ws = wb.create_sheet("CATÁLOGO", 0)
    ws["A1"] = "CATÁLOGO DE CATEGORÍAS DE GASTO · la misma lista en el Excel y en el sistema"
    ws["A1"].font = Font(bold=True, size=14, color=OSCURO)
    ws["A2"] = "Solo Fijo y Variable entran al punto de equilibrio. Financiamiento se ve aparte (PE incluyendo deudas). No inventar categorías nuevas: si algo no encaja, avisar a Jahnn."
    ws["A2"].font = Font(italic=True, color="5C5C5C")
    hdr = ["Categoría", "Tipo", "¿Entra al punto de equilibrio?", "Qué va aquí", "Ejemplos"]
    for i, h in enumerate(hdr, 1):
        c = ws.cell(4, i, h); c.font = H; c.fill = HFILL; c.alignment = Alignment(vertical="center")
    for j, cat in enumerate(CATS):
        r = 5 + j
        entra = "Sí" if cat["tipo"] in ("Fijo", "Variable") else ("Solo en 'PE incluyendo deudas'" if cat["tipo"] == "Financiamiento" else "No")
        vals = [cat["nombre"], cat["tipo"], entra, cat["descripcion"], cat["ejemplos"]]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v); c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="top")
            c.fill = PatternFill("solid", fgColor=TIPO_FILL[cat["tipo"]])
        ws.cell(r, 1).font = Font(bold=True)
    for col, w in zip("ABCDE", [28, 15, 22, 48, 70]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    n = len(CATS)
    wb.defined_names["LISTA_CATEGORIAS"] = DefinedName("LISTA_CATEGORIAS", attr_text=f"'CATÁLOGO'!$A$5:$A${4+n}")
    wb.defined_names["CATALOGO_TABLA"] = DefinedName("CATALOGO_TABLA", attr_text=f"'CATÁLOGO'!$A$5:$B${4+n}")
    return ws

def hoja_reglas(wb):
    ws = wb.create_sheet("REGLAS", 1)
    hdr = ["Palabra clave", "Categoría", "Cómo se lee (· = espacio)"]
    for i, h in enumerate(hdr, 1):
        c = ws.cell(1, i, h); c.font = H; c.fill = HFILL
    for j, rg in enumerate(REGLAS):
        r = 2 + j
        ws.cell(r, 1, rg["clave"]); ws.cell(r, 2, rg["categoria"])
        ws.cell(r, 3, f'=SUBSTITUTE(A{r}," ","·")').font = Font(color="5C5C5C")
    ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 28; ws.column_dimensions["C"].width = 34
    ws.column_dimensions["E"].width = 70
    notas = [
        "CÓMO FUNCIONA",
        "El Excel busca cada palabra clave dentro del Concepto y del Proveedor de cada gasto.",
        "Si varias coinciden, MANDA LA QUE ESTÁ MÁS ABAJO en esta lista.",
        "Por eso arriba están los proveedores (lo más general) y abajo las excepciones",
        "(ej. 'BOLSA DE HIELO' es insumo aunque diga BOLSA).",
        "",
        "PARA ENSEÑARLE ALGO NUEVO",
        "Agrega una fila AL FINAL (sin dejar filas vacías en medio):",
        "palabra en MAYÚSCULAS y sin tildes + la categoría exacta del CATÁLOGO.",
        "Las palabras cortas llevan espacios alrededor para no encontrarse dentro de otras",
        "(' IR ' no encuentra 'GIRO').",
    ]
    for k, t in enumerate(notas):
        c = ws.cell(2 + k, 5, t)
        if k in (0, 6): c.font = Font(bold=True, color=OSCURO)
    ws.freeze_panes = "A2"
    dv = DataValidation(type="list", formula1="=LISTA_CATEGORIAS", allow_blank=True, errorStyle="stop",
                        error="Usa una categoría del CATÁLOGO.", errorTitle="Categoría")
    ws.add_data_validation(dv); dv.add("B2:B2000")
    wb.defined_names["REGLAS_CLAVE"] = DefinedName("REGLAS_CLAVE", attr_text="OFFSET(REGLAS!$A$2,0,0,COUNTA(REGLAS!$A:$A)-1,1)")
    wb.defined_names["REGLAS_CATEGORIA"] = DefinedName("REGLAS_CATEGORIA", attr_text="OFFSET(REGLAS!$B$2,0,0,COUNTA(REGLAS!$A:$A)-1,1)")

def hoja_como_usar(wb, hoja_mes, hoja_pe):
    ws = wb.create_sheet("CÓMO USAR", 0)
    ws.column_dimensions["A"].width = 4; ws.column_dimensions["B"].width = 110
    ws["B1"] = "CÓMO SE CLASIFICAN LOS GASTOS DESDE AHORA"; ws["B1"].font = Font(bold=True, size=15, color=OSCURO)
    pasos = [
        ("1. Registra como siempre", "Fecha, Ing./Gsto., Proveedor, Concepto y monto. Escribe el concepto claro: 'PECHUGAS DE POLLO', no 'VARIOS'."),
        ("2. El Grupo se pone solo", "En los gastos (G), la columna C 'Grupo' se llena sola con la categoría del CATÁLOGO, leyendo el Concepto y el Proveedor. La columna W dice si es Fijo, Variable, etc."),
        ("3. Si está mal o dice POR ACLARAR (en rojo)", "Elige la categoría correcta en la columna V 'Corregir' (lista desplegable). Esa manda sobre la automática. No escribas encima de la columna C."),
        ("4. Para que aprenda", "Si un gasto se repite y siempre sale mal, agrega una fila al final de la pestaña REGLAS (palabra clave + categoría). Desde ahí se clasifica solo."),
        ("5. Ingresos (I)", "En los ingresos escribe el Grupo en la columna C como siempre (Ventas, SALDO, OTROS…)."),
        ("6. Mes nuevo", f"Duplica la pestaña '{hoja_mes}' (las fórmulas de las columnas C, U, W y X vienen solas) y la pestaña '{hoja_pe}'. En la PE nueva cambia solo las 2 celdas amarillas (nombre de la hoja de gastos y de ventas). Agrega el mes en 'PE Resumen'."),
        ("7. Si la columna C sale vacía en un gasto", "Revisa que la columna B diga G. Si agregaste filas nuevas al final, copia hacia abajo la fila de arriba (columnas C, U, W y X)."),
    ]
    for i, (t, d) in enumerate(pasos):
        r = 3 + i * 2
        ws.cell(r, 2, t).font = Font(bold=True, color=VERDE, size=12)
        c = ws.cell(r + 1, 2, d); c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r + 1].height = 32
    r = 3 + len(pasos) * 2 + 1
    ws.cell(r, 2, "Por qué: con una sola lista, el Excel y el sistema dicen lo mismo, y el punto de equilibrio se calcula con los mismos criterios todos los meses. Lo que nadie sabe qué es queda POR ACLARAR y no ensucia el cálculo.").alignment = Alignment(wrap_text=True)
    ws.row_dimensions[r].height = 32
    ws.sheet_properties.tabColor = VERDE

def convertir_mes(wb, hoja, overrides):
    ws = wb[hoja]
    lastM = max(r for r in range(1, ws.max_row + 1) if isinstance(ws.cell(r, 13).value, str) and str(ws.cell(r, 13).value).startswith("="))
    lastData = max(r for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value not in (None, ""))
    last = max(lastM, lastData)
    ref = ws["C3"]
    for col, t, w in [(21, "Categoría sugerida (automática)", 26), (22, "Corregir (opcional)", 26), (23, "Tipo", 16), (24, "texto (no tocar)", 10)]:
        c = ws.cell(3, col, t)
        c.font = copy.copy(ref.font); c.fill = copy.copy(ref.fill); c.border = copy.copy(ref.border); c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.cell(3, 22).fill = PatternFill("solid", fgColor="FFF2CC")
    ws.column_dimensions["X"].hidden = True
    cambios = 0
    for r in range(4, last + 1):
        b = ws.cell(r, 2).value
        b = str(b).strip().upper() if b is not None else ""
        ws.cell(r, 24, "=\" \"&" + acentos(f"UPPER(TRIM($F{r}&\" \"&$D{r}))") + "&\" \"")
        ws.cell(r, 21, f'=IF($B{r}="G",IFERROR(LOOKUP(2^15,SEARCH(REGLAS_CLAVE,$X{r}),REGLAS_CATEGORIA),"POR ACLARAR"),"")')
        ws.cell(r, 23, f'=IF($C{r}="","",IFERROR(VLOOKUP($C{r},CATALOGO_TABLA,2,FALSE),IF($B{r}="G","⚠ no está en el CATÁLOGO","")))')
        if b == "I":
            continue  # ingresos: Kelly escribe el Grupo en C como siempre
        if r in overrides:
            ws.cell(r, 22, overrides[r])
        ws.cell(r, 3, f'=IF($V{r}<>"",$V{r},IF($B{r}="G",$U{r},""))')
        cambios += 1
    dv = DataValidation(type="list", formula1="=LISTA_CATEGORIAS", allow_blank=True, errorStyle="stop",
                        error="Elige una categoría del CATÁLOGO.", errorTitle="Corregir")
    ws.add_data_validation(dv); dv.add(f"V5:V{last}")
    ws.conditional_formatting.add(f"C5:C{last}", FormulaRule(formula=['$C5="POR ACLARAR"'], fill=PatternFill("solid", fgColor="F8D7DA"), font=Font(bold=True, color="9C0006")))
    ws.conditional_formatting.add(f"C5:C{last}", FormulaRule(formula=['$V5<>""'], fill=PatternFill("solid", fgColor="DDEBF7")))
    ws.conditional_formatting.add(f"W5:W{last}", FormulaRule(formula=['LEFT($W5,1)="⚠"'], fill=PatternFill("solid", fgColor="FFE699")))
    return last, cambios

def hoja_pe(wb, hoja_mes, hoja_vtas, titulo, mes_etq):
    nombre = f"PE {mes_etq}"
    idx = wb.sheetnames.index(nombre) if nombre in wb.sheetnames else len(wb.sheetnames)
    if nombre in wb.sheetnames:
        del wb[nombre]
    ws = wb.create_sheet(nombre, idx)
    ws.column_dimensions["A"].width = 2; ws.column_dimensions["B"].width = 52; ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 18; ws.column_dimensions["E"].width = 2; ws.column_dimensions["F"].width = 80
    ws["B1"] = f"PUNTO DE EQUILIBRIO — {mes_etq} · {titulo} (automático, lista única de categorías)"; ws["B1"].font = Font(bold=True, size=13, color=OSCURO)
    ws["B2"] = "Suma los gastos 'G' por la categoría de la columna C y la clasifica con el CATÁLOGO. Solo Fijo y Variable entran al PE."
    ws["B2"].font = Font(italic=True, color="5C5C5C")
    amarillo = PatternFill("solid", fgColor="FFF2CC")
    ws["B3"] = "Hoja de gastos del mes:"; ws["D3"] = hoja_mes; ws["D3"].fill = amarillo
    ws["B4"] = "Hoja de ventas del mes:"; ws["D4"] = hoja_vtas; ws["D4"].fill = amarillo
    ws["F3"] = "← Mes nuevo: cambia solo estas 2 celdas amarillas."
    G = lambda col: f'INDIRECT("\'"&$D$3&"\'!${col}$5:${col}$3000")'
    def sec(r, t):
        ws.cell(r, 2, t).font = Font(bold=True, color=VERDE, size=12)
    sec(6, "1. VENTAS DEL MES")
    ws["B7"] = "Ventas totales BITE"; ws["D7"] = f'=INDIRECT("\'"&$D$4&"\'!E194")'; ws["F7"] = "Total de Control de VTAS, celda E194"
    ws["B8"] = "Días con venta"; ws["D8"] = f'=COUNT(INDIRECT("\'"&$D$4&"\'!B2:B187"))'
    ws["B9"] = "Venta promedio diaria"; ws["D9"] = "=IF(D8=0,0,D7/D8)"
    sec(11, "2. GASTOS POR CATEGORÍA")
    for i, h in enumerate(["Categoría", "Tipo", "Importe del mes"]):
        c = ws.cell(12, 2 + i, h); c.font = H; c.fill = HFILL
    r0 = 13
    for j, cat in enumerate(CATS):
        r = r0 + j
        ws.cell(r, 2, cat["nombre"]); ws.cell(r, 3, f"=VLOOKUP(B{r},CATALOGO_TABLA,2,FALSE)")
        ws.cell(r, 4, f'=SUMIFS({G("K")},{G("C")},B{r},{G("B")},"G")+SUMIFS({G("L")},{G("C")},B{r},{G("B")},"G")')
        ws.cell(r, 2).fill = ws.cell(r, 3).fill = PatternFill("solid", fgColor=TIPO_FILL[cat["tipo"]])
    rN = r0 + len(CATS) - 1
    rng = lambda c: f"{c}{r0}:{c}{rN}"
    porAcl = r0 + [c["nombre"] for c in CATS].index("POR ACLARAR")
    sec(rN + 2, "3. TOTALES POR TIPO")
    t = rN + 3
    filas = [
        ("Total Costos Variables", f'=SUMIFS({rng("D")},{rng("C")},"Variable")'),
        ("Total Costos Fijos", f'=SUMIFS({rng("D")},{rng("C")},"Fijo")'),
        ("Total Financiamiento (préstamos y tarjetas)", f'=SUMIFS({rng("D")},{rng("C")},"Financiamiento")'),
        ("Total Inversión (equipos y remodelación)", f'=SUMIFS({rng("D")},{rng("C")},"Inversión")'),
        ("Total No es gasto (utilidades, ahorro, entre sedes, devoluciones, por aclarar)", f'=SUMIFS({rng("D")},{rng("C")},"No es gasto")'),
        ("   de los cuales POR ACLARAR (revisar)", f"=D{porAcl}"),
    ]
    for k, (lab, f) in enumerate(filas):
        ws.cell(t + k, 2, lab); ws.cell(t + k, 4, f)
    tv, tf, tfin = t, t + 1, t + 2
    p = t + len(filas) + 1
    sec(p, "4. PUNTO DE EQUILIBRIO")
    pe = [
        ("Margen de contribución (%)", f"=IF(D7=0,0,(D7-D{tv})/D7)", "0.0%"),
        ("PUNTO DE EQUILIBRIO MENSUAL (S/)", f"=IF(D{p+1}<=0,0,D{tf}/D{p+1})", None),
        ("Punto de equilibrio incluyendo deudas (S/)", f"=IF(D{p+1}<=0,0,(D{tf}+D{tfin})/D{p+1})", None),
        ("Punto de equilibrio diario (S/)", f"=IF(D8=0,0,D{p+2}/D8)", None),
        ("Ventas netas del mes", "=D7", None),
        ("Margen de seguridad (S/)", f"=D{p+5}-D{p+2}", None),
        ("UTILIDAD OPERATIVA DEL MES", f"=D7-D{tv}-D{tf}", None),
        ("Estado", f'=IF(D{p+5}>=D{p+2},"✔ SOBRE EL EQUILIBRIO","✖ DEBAJO DEL EQUILIBRIO")', None),
    ]
    for k, (lab, f, fmt) in enumerate(pe):
        ws.cell(p + 1 + k, 2, lab); c = ws.cell(p + 1 + k, 4, f)
        if fmt: c.number_format = fmt
    ws.cell(p + 2, 2).font = Font(bold=True); ws.cell(p + 2, 4).font = Font(bold=True, color=VERDE)
    ws.cell(p + 3, 6, "Lo que hay que vender para pagar también las cuotas de préstamos y tarjetas.")
    c = p + len(pe) + 2
    sec(c, "5. CONCILIACIÓN")
    ws.cell(c + 1, 2, 'Total de gastos "G" registrados en el mes')
    ws.cell(c + 1, 4, f'=SUMIFS({G("K")},{G("B")},"G")+SUMIFS({G("L")},{G("B")},"G")')
    ws.cell(c + 2, 2, "Total clasificado (todas las categorías)"); ws.cell(c + 2, 4, f"=SUM({rng('D')})")
    ws.cell(c + 3, 2, 'Debe ser "0"'); ws.cell(c + 3, 4, f"=ROUND(D{c+1}-D{c+2},2)")
    ws.cell(c + 3, 6, "Si no es 0: hay un gasto cuyo Grupo no está en el CATÁLOGO (búscalo en la columna W del mes: dice ⚠).")
    for row in ws.iter_rows(min_row=7, max_row=c + 3, min_col=4, max_col=4):
        for cell in row:
            if cell.number_format == "General": cell.number_format = '#,##0.00'
    return {"ventas": 7, "var": tv, "fijos": tf, "fin": tfin, "pe": p + 2, "pe_deuda": p + 3, "util": p + 7, "conc": c + 3}

def actualizar_resumen(wb, mes_etq, cel):
    if "PE Resumen" not in wb.sheetnames: return None
    ws = wb["PE Resumen"]
    fila = next((r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 2).value).strip() == mes_etq), None)
    if fila is None: return None
    n = f"'PE {mes_etq}'"
    ws.cell(fila, 3, f"={n}!D{cel['ventas']}"); ws.cell(fila, 4, f"={n}!D{cel['var']}"); ws.cell(fila, 5, f"={n}!D{cel['fijos']}")
    ws.cell(fila, 6, f"={n}!D{cel['pe']}"); ws.cell(fila, 7, f"={n}!D{cel['util']}")
    ws.cell(fila, 9, f"={n}!D{cel['conc']}")
    hdr = ws.cell(4, 10, "PE incluyendo deudas"); hdr.font = copy.copy(ws.cell(4, 9).font); hdr.fill = copy.copy(ws.cell(4, 9).fill)
    ws.cell(fila, 10, f"={n}!D{cel['pe_deuda']}").number_format = ws.cell(fila, 6).number_format
    ws.cell(fila + 2 if ws.cell(fila + 1, 2).value is None else fila + 3, 2,
            f"Desde {mes_etq} con la lista única de categorías (CATÁLOGO); los meses anteriores usan 'Categorías PE'.").font = Font(italic=True, color="5C5C5C")
    return fila

for sede, (f, hoja, vtas, titulo, ov) in SEDES.items():
    wb = openpyxl.load_workbook(D + f)
    last, cambios = convertir_mes(wb, hoja, ov)
    hoja_catalogo(wb); hoja_reglas(wb); hoja_como_usar(wb, hoja, "PE SEP26")
    cel = hoja_pe(wb, hoja, vtas, titulo, "SEP26")
    fila = actualizar_resumen(wb, "SEP26", cel)
    wb.active = 0
    out = OUT + f"INGRESOS_GASTOS_SUMMARY_{sede}_19_09_2026_CATEGORIAS.xlsx"
    wb.save(out)
    print(sede, "filas hasta", last, "gastos con fórmula", cambios, "resumen fila", fila, "->", out)
